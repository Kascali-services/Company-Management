"""
Integration tests for Export routes.

Tests CSV and Excel export functionality for the /api/export endpoints.
"""

import pytest
import csv
import io
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.test_helpers import assert_response_status


@pytest.mark.integration
class TestExportCSVEndpoint:
    """Test suite for GET /api/export/csv endpoint."""

    def test_export_csv_empty_database(self, client: TestClient):
        """
        Test CSV export with empty database.

        Arrange: Empty database
        Act: GET /api/export/csv
        Assert: Returns CSV with headers only
        """
        # Arrange & Act
        response = client.get("/api/export/csv")

        # Assert
        assert_response_status(response, 200)
        assert response.headers["content-type"] == "text/csv; charset=utf-8"
        assert "attachment" in response.headers["content-disposition"]
        assert "unternehmen.csv" in response.headers["content-disposition"]

        # Parse CSV content
        csv_content = response.text
        reader = csv.reader(io.StringIO(csv_content))
        rows = list(reader)

        # Should have header row only
        assert len(rows) == 1
        header = rows[0]
        assert "Unternehmen" in header
        assert "Bundesland" in header
        assert "Ansprechpartner Vorname" in header

    def test_export_csv_with_companies(self, client: TestClient, create_unternehmen):
        """
        Test CSV export with companies.

        Arrange: Create multiple companies
        Act: GET /api/export/csv
        Assert: CSV contains all companies
        """
        # Arrange
        create_unternehmen(name="Firma A", bundesland="Bayern", stadt="München", plz="80331")
        create_unternehmen(name="Firma B", bundesland="Berlin", stadt="Berlin", plz="10115")

        # Act
        response = client.get("/api/export/csv")

        # Assert
        assert_response_status(response, 200)

        # Parse CSV content
        csv_content = response.text
        reader = csv.reader(io.StringIO(csv_content))
        rows = list(reader)

        # Should have header + 2 data rows
        assert len(rows) == 3

        # Check data rows
        firma_names = [row[0] for row in rows[1:]]
        assert "Firma A" in firma_names
        assert "Firma B" in firma_names

    def test_export_csv_with_ansprechpartner(self, client: TestClient, create_unternehmen, create_person, test_db):
        """
        Test CSV export includes Ansprechpartner information.

        Arrange: Create company with Ansprechpartner
        Act: GET /api/export/csv
        Assert: CSV contains Ansprechpartner details
        """
        # Arrange
        unternehmen = create_unternehmen(name="Test GmbH")
        person = create_person(
            firma_id=unternehmen.id,
            rolle="Ansprechpartner",
            vorname="Anna",
            nachname="Kontakt",
            email="anna@test.de"
        )
        unternehmen.ansprechpartner_id = person.id
        test_db.commit()

        # Act
        response = client.get("/api/export/csv")

        # Assert
        assert_response_status(response, 200)

        # Parse CSV
        csv_content = response.text
        reader = csv.reader(io.StringIO(csv_content))
        rows = list(reader)

        # Find the company row
        company_row = rows[1]
        assert company_row[0] == "Test GmbH"
        assert "Anna" in company_row  # Ansprechpartner Vorname
        assert "Kontakt" in company_row  # Ansprechpartner Nachname
        assert "anna@test.de" in company_row  # Ansprechpartner Email

    def test_export_csv_with_empfehler_count(self, client: TestClient, create_unternehmen, create_person):
        """
        Test CSV export includes Empfehler count.

        Arrange: Create company with multiple Empfehler
        Act: GET /api/export/csv
        Assert: CSV shows correct Empfehler count
        """
        # Arrange
        unternehmen = create_unternehmen(name="Test GmbH")
        for i in range(3):
            create_person(
                firma_id=unternehmen.id,
                rolle="Empfehler",
                vorname=f"Emp{i}"
            )

        # Act
        response = client.get("/api/export/csv")

        # Assert
        assert_response_status(response, 200)

        # Parse CSV
        csv_content = response.text
        reader = csv.reader(io.StringIO(csv_content))
        rows = list(reader)

        # Check Empfehler count (last column)
        company_row = rows[1]
        empfehler_count = company_row[-1]
        assert empfehler_count == "3"

    def test_export_csv_without_ansprechpartner(self, client: TestClient, create_unternehmen):
        """
        Test CSV export for company without Ansprechpartner.

        Arrange: Create company without Ansprechpartner
        Act: GET /api/export/csv
        Assert: CSV has empty Ansprechpartner fields
        """
        # Arrange
        create_unternehmen(name="Solo GmbH")

        # Act
        response = client.get("/api/export/csv")

        # Assert
        assert_response_status(response, 200)

        # Parse CSV
        csv_content = response.text
        reader = csv.reader(io.StringIO(csv_content))
        rows = list(reader)

        company_row = rows[1]
        # Ansprechpartner fields should be empty
        assert company_row[6] == ""  # Ansprechpartner Vorname
        assert company_row[7] == ""  # Ansprechpartner Nachname
        assert company_row[8] == ""  # Ansprechpartner Email

    def test_export_csv_special_characters(self, client: TestClient, create_unternehmen, create_person, test_db):
        """
        Test CSV export handles special characters correctly.

        Arrange: Create company with umlauts and special chars
        Act: GET /api/export/csv
        Assert: CSV preserves special characters
        """
        # Arrange
        unternehmen = create_unternehmen(
            name="Müller & Söhne GmbH",
            bundesland="Baden-Württemberg",
            stadt="Köln"
        )
        person = create_person(
            firma_id=unternehmen.id,
            rolle="Ansprechpartner",
            vorname="Jürgen",
            nachname="Größmann"
        )
        unternehmen.ansprechpartner_id = person.id
        test_db.commit()

        # Act
        response = client.get("/api/export/csv")

        # Assert
        assert_response_status(response, 200)

        csv_content = response.text
        assert "Müller & Söhne GmbH" in csv_content
        assert "Baden-Württemberg" in csv_content
        assert "Jürgen" in csv_content
        assert "Größmann" in csv_content

    def test_export_csv_headers_correct_order(self, client: TestClient):
        """
        Test CSV headers are in correct order.

        Arrange: Empty database
        Act: GET /api/export/csv
        Assert: Headers match expected order
        """
        # Arrange & Act
        response = client.get("/api/export/csv")

        # Assert
        csv_content = response.text
        reader = csv.reader(io.StringIO(csv_content))
        headers = next(reader)

        expected_headers = [
            'Unternehmen', 'Bundesland', 'Stadt', 'PLZ', 'Straße', 'Hausnummer',
            'Ansprechpartner Vorname', 'Ansprechpartner Nachname', 'Ansprechpartner Email',
            'Anzahl Empfehler'
        ]

        assert headers == expected_headers

    def test_export_csv_multiple_companies(self, client: TestClient, multiple_unternehmen):
        """
        Test CSV export with multiple companies.

        Arrange: Create 5 companies via fixture
        Act: GET /api/export/csv
        Assert: All companies in CSV
        """
        # Arrange - fixture creates 5 companies

        # Act
        response = client.get("/api/export/csv")

        # Assert
        assert_response_status(response, 200)

        csv_content = response.text
        reader = csv.reader(io.StringIO(csv_content))
        rows = list(reader)

        # Header + 5 data rows
        assert len(rows) == 6


@pytest.mark.integration
class TestExportExcelEndpoint:
    """Test suite for GET /api/export/excel endpoint."""

    def test_export_excel_empty_database(self, client: TestClient):
        """
        Test Excel export with empty database.

        Arrange: Empty database
        Act: GET /api/export/excel
        Assert: Returns Excel file with headers only
        """
        # Arrange & Act
        response = client.get("/api/export/excel")

        # Assert
        assert_response_status(response, 200)
        assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert "attachment" in response.headers["content-disposition"]
        assert "unternehmen.xlsx" in response.headers["content-disposition"]

        # Parse Excel content
        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active
        assert ws.title == "Unternehmen"

        # Should have header row only
        assert ws.max_row == 1
        headers = [cell.value for cell in ws[1]]
        assert "Unternehmen" in headers
        assert "Bundesland" in headers

    def test_export_excel_with_companies(self, client: TestClient, create_unternehmen):
        """
        Test Excel export with companies.

        Arrange: Create multiple companies
        Act: GET /api/export/excel
        Assert: Excel contains all companies
        """
        # Arrange
        create_unternehmen(name="Firma A", bundesland="Bayern")
        create_unternehmen(name="Firma B", bundesland="Berlin")

        # Act
        response = client.get("/api/export/excel")

        # Assert
        assert_response_status(response, 200)

        # Parse Excel
        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active

        # Header + 2 data rows
        assert ws.max_row == 3

        # Check company names
        names = [ws.cell(row=i, column=1).value for i in range(2, 4)]
        assert "Firma A" in names
        assert "Firma B" in names

    def test_export_excel_with_ansprechpartner(self, client: TestClient, create_unternehmen, create_person, test_db):
        """
        Test Excel export includes Ansprechpartner information.

        Arrange: Create company with Ansprechpartner
        Act: GET /api/export/excel
        Assert: Excel contains Ansprechpartner details
        """
        # Arrange
        unternehmen = create_unternehmen(name="Test GmbH")
        person = create_person(
            firma_id=unternehmen.id,
            rolle="Ansprechpartner",
            vorname="Anna",
            nachname="Kontakt",
            email="anna@test.de"
        )
        unternehmen.ansprechpartner_id = person.id
        test_db.commit()

        # Act
        response = client.get("/api/export/excel")

        # Assert
        assert_response_status(response, 200)

        # Parse Excel
        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active

        # Check Ansprechpartner data (columns 7, 8, 9)
        assert ws.cell(row=2, column=7).value == "Anna"
        assert ws.cell(row=2, column=8).value == "Kontakt"
        assert ws.cell(row=2, column=9).value == "anna@test.de"

    def test_export_excel_with_empfehler_count(self, client: TestClient, create_unternehmen, create_person):
        """
        Test Excel export includes Empfehler count.

        Arrange: Create company with multiple Empfehler
        Act: GET /api/export/excel
        Assert: Excel shows correct Empfehler count
        """
        # Arrange
        unternehmen = create_unternehmen(name="Test GmbH")
        for i in range(3):
            create_person(
                firma_id=unternehmen.id,
                rolle="Empfehler",
                vorname=f"Emp{i}"
            )

        # Act
        response = client.get("/api/export/excel")

        # Assert
        assert_response_status(response, 200)

        # Parse Excel
        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active

        # Check Empfehler count (last column, column 10)
        assert ws.cell(row=2, column=10).value == 3

    def test_export_excel_without_ansprechpartner(self, client: TestClient, create_unternehmen):
        """
        Test Excel export for company without Ansprechpartner.

        Arrange: Create company without Ansprechpartner
        Act: GET /api/export/excel
        Assert: Excel has empty Ansprechpartner fields
        """
        # Arrange
        create_unternehmen(name="Solo GmbH")

        # Act
        response = client.get("/api/export/excel")

        # Assert
        assert_response_status(response, 200)

        # Parse Excel
        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active

        # Ansprechpartner fields should be empty
        assert ws.cell(row=2, column=7).value == ""
        assert ws.cell(row=2, column=8).value == ""
        assert ws.cell(row=2, column=9).value == ""

    def test_export_excel_special_characters(self, client: TestClient, create_unternehmen, create_person, test_db):
        """
        Test Excel export handles special characters correctly.

        Arrange: Create company with umlauts and special chars
        Act: GET /api/export/excel
        Assert: Excel preserves special characters
        """
        # Arrange
        unternehmen = create_unternehmen(
            name="Müller & Söhne GmbH",
            bundesland="Baden-Württemberg",
            stadt="Köln"
        )
        person = create_person(
            firma_id=unternehmen.id,
            rolle="Ansprechpartner",
            vorname="Jürgen",
            nachname="Größmann"
        )
        unternehmen.ansprechpartner_id = person.id
        test_db.commit()

        # Act
        response = client.get("/api/export/excel")

        # Assert
        assert_response_status(response, 200)

        # Parse Excel
        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active

        assert ws.cell(row=2, column=1).value == "Müller & Söhne GmbH"
        assert ws.cell(row=2, column=2).value == "Baden-Württemberg"
        assert ws.cell(row=2, column=3).value == "Köln"
        assert ws.cell(row=2, column=7).value == "Jürgen"
        assert ws.cell(row=2, column=8).value == "Größmann"

    def test_export_excel_headers_correct(self, client: TestClient):
        """
        Test Excel headers are correct.

        Arrange: Empty database
        Act: GET /api/export/excel
        Assert: Headers match expected
        """
        # Arrange & Act
        response = client.get("/api/export/excel")

        # Assert
        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active

        expected_headers = [
            'Unternehmen', 'Bundesland', 'Stadt', 'PLZ', 'Straße', 'Hausnummer',
            'Ansprechpartner Vorname', 'Ansprechpartner Nachname', 'Ansprechpartner Email',
            'Anzahl Empfehler'
        ]

        actual_headers = [ws.cell(row=1, column=i).value for i in range(1, 11)]
        assert actual_headers == expected_headers

    def test_export_excel_multiple_companies(self, client: TestClient, multiple_unternehmen):
        """
        Test Excel export with multiple companies.

        Arrange: Create 5 companies via fixture
        Act: GET /api/export/excel
        Assert: All companies in Excel
        """
        # Arrange - fixture creates 5 companies

        # Act
        response = client.get("/api/export/excel")

        # Assert
        assert_response_status(response, 200)

        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active

        # Header + 5 data rows
        assert ws.max_row == 6

    def test_export_excel_complete_scenario(self, client: TestClient, unternehmen_with_persons):
        """
        Test Excel export with complete company scenario.

        Arrange: Company with Ansprechpartner and multiple Empfehler
        Act: GET /api/export/excel
        Assert: All data correctly exported
        """
        # Arrange - fixture creates complete scenario
        data = unternehmen_with_persons

        # Act
        response = client.get("/api/export/excel")

        # Assert
        assert_response_status(response, 200)

        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active

        # Check Ansprechpartner
        assert ws.cell(row=2, column=7).value == data["ansprechpartner"].vorname

        # Check Empfehler count (3 Empfehler)
        assert ws.cell(row=2, column=10).value == 3
